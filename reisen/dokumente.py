"""Excel-Export (mit echten Formeln) und PDF-Reiseübersicht."""
from datetime import timedelta
from io import BytesIO

from django.conf import settings
from django.db.models import Prefetch
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from kalender.tage import zeitraum_text

from . import mitreise
from .models import Buchung, Buchungsstatus, Etappe, Kostenart, Reise

EURO = '#,##0.00 "€"'
ZAHL2 = "#,##0.00"
KM = '#,##0 "km"'
DATUM = "DD.MM.YYYY"
WOTAG_DATUM = "DDD DD.MM.YYYY"
UHR = "HH:MM"
DAUER = "[H]:MM"

BLAU = "173F73"
KOPF = PatternFill("solid", fgColor=BLAU)
KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")
FETT = Font(bold=True)
LINIE_OBEN = Border(top=Side(style="medium", color=BLAU))


def reise_laden(pk, familie, mitreisend=True):
    """Reise mit allem Zubehör, so wie die Familie sie sieht – eigene und (mit `mitreisend`)
    gemeinsame Reisen; sonst Reise.DoesNotExist."""
    reisen = (Reise.objects.select_related("fahrzeugprofil", "familie")
              .prefetch_related(
                  Prefetch("etappen", queryset=Etappe.objects.select_related(
                      "von_ort", "nach_ort", "fahrzeugprofil")),
                  Prefetch("buchungen", queryset=Buchung.objects.select_related("ort")),
                  "kosten", "links", mitreise.prefetch_familie()))
    reisen = mitreise.sichtbar(familie, reisen) if mitreisend else reisen.filter(familie=familie)
    return mitreise.fuer([reisen.get(pk=pk)], familie)[0]


def stand_jetzt():
    """Zeitpunkt der Erzeugung (Ortszeit) – steht im Dokument und im Dateinamen."""
    return timezone.localtime()


def stempel(stand):
    return f"Stand_{stand:%Y-%m-%d_%H%M}"


def dateiname(reise, endung, stand):
    titel = "".join(c if c.isalnum() else "_" for c in reise.titel).strip("_")
    return f"{reise.urlaub_beginn:%Y-%m-%d}_{titel}_{stempel(stand)}.{endung}"


# --- Excel -------------------------------------------------------------------


def _kopfzeile(ws, spalten, breiten):
    ws.append(spalten)
    for i, breite in enumerate(breiten, start=1):
        zelle = ws.cell(row=1, column=i)
        zelle.fill, zelle.font = KOPF, KOPF_SCHRIFT
        zelle.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"


def _formate(ws, zeilen, formate):
    """formate: {Spaltenbuchstabe: Zahlenformat} für Zeilen 2..zeilen+1"""
    for spalte, fmt in formate.items():
        for z in range(2, zeilen + 2):
            ws[f"{spalte}{z}"].number_format = fmt


def _bereich(spalte, anzahl):
    """B2:B{n+1}; bei 0 Zeilen eine leere Zeile, damit SUMME gültig bleibt."""
    return f"{spalte}2:{spalte}{max(anzahl, 1) + 1}"


def excel(reise, stand):
    wb = Workbook()
    ue = wb.active
    ue.title = "Übersicht"

    # Etappen
    et = wb.create_sheet("Etappen")
    _kopfzeile(et, ["Nr.", "Datum", "Von", "Nach", "Start", "Fahrzeit", "Pause", "Ankunft",
                    "Entfernung", "L/100 km", "Liter", "Länder", "Vignette", "Fähre", "Karte"],
               [5, 16, 26, 26, 8, 9, 8, 9, 11, 9, 9, 18, 10, 7, 10])
    etappen = list(reise.etappen.all())
    for i, e in enumerate(etappen, start=2):
        d = e.routing_details or {}
        et.append([i - 1, e.datum, e.von_ort.name, e.nach_ort.name, e.startzeit,
                   e.fahrzeit, e.pause or timedelta(),
                   f"=IF(OR(E{i}=\"\",F{i}=\"\"),\"\",E{i}+F{i}+G{i})",
                   e.entfernung_km,
                   # TomTom-Liter (Verbrauchskurve) als Wert, sonst Formel km × L/100 km
                   *(("TomTom", e.tomtom_liter) if e.tomtom_liter is not None else
                     (e.verbrauch, f"=IF(OR(I{i}=\"\",J{i}=\"\"),0,I{i}*J{i}/100)")),
                   " → ".join(d.get("laender", [])), ", ".join(d.get("vignetten", [])),
                   "ja" if d.get("faehre") else "", "Route" if e.karten_link else ""])
        if e.karten_link:
            et[f"O{i}"].hyperlink = e.karten_link
            et[f"O{i}"].style = "Hyperlink"
    _formate(et, len(etappen), {"B": WOTAG_DATUM, "E": UHR, "F": DAUER, "G": DAUER, "H": UHR,
                                "I": KM, "J": "0.0", "K": "0.0"})
    n_et = len(etappen)
    s = n_et + 2
    et[f"A{s}"], et[f"F{s}"], et[f"I{s}"], et[f"K{s}"] = (
        "Summe", f"=SUM({_bereich('F', n_et)})", f"=SUM({_bereich('I', n_et)})", f"=SUM({_bereich('K', n_et)})")
    for sp, fmt in (("F", DAUER), ("I", KM), ("K", "0.0")):
        et[f"{sp}{s}"].number_format = fmt
    for sp in "ABCDEFGHIJKLMNO":
        et[f"{sp}{s}"].font, et[f"{sp}{s}"].border = FETT, LINIE_OBEN

    # Buchungen
    bu = wb.create_sheet("Buchungen")
    _kopfzeile(bu, ["Ort", "Anreise", "Abreise", "Nächte", "Preis/Nacht", "Gesamtpreis", "Summe", "Währung",
                    "Kurs (1 € =)", "Summe €", "Anzahlung", "Rest", "Rest €", "Rest fällig", "Status",
                    "Buchungsnr.", "Stellplatz", "Website"],
               [30, 12, 12, 8, 11, 10, 11, 8, 10, 12, 11, 11, 12, 12, 11, 16, 12, 10])
    buchungen = list(reise.buchungen.all())
    for i, b in enumerate(buchungen, start=2):
        # Summe: Gesamtpreis, wenn eingetragen, sonst Nächte × Preis/Nacht
        bu.append([b.ort.name, b.von, b.bis, f"=C{i}-B{i}", b.preis_pro_nacht, b.gesamtpreis,
                   f'=IF(F{i}<>"",F{i},D{i}*E{i})',
                   b.waehrung, b.wechselkurs, f"=ROUND(G{i}/I{i},2)", b.anzahlung, f"=G{i}-K{i}",
                   f"=ROUND(L{i}/I{i},2)", b.rest_faellig_am, b.get_status_display(), b.buchungsnummer,
                   b.stellplatz, "Website" if b.ort.website else ""])
        if b.ort.website:
            bu[f"R{i}"].hyperlink = b.ort.website
            bu[f"R{i}"].style = "Hyperlink"
    _formate(bu, len(buchungen), {"B": DATUM, "C": DATUM, "E": ZAHL2, "F": ZAHL2, "G": ZAHL2, "I": "0.0000",
                                  "J": EURO, "K": ZAHL2, "L": ZAHL2, "M": EURO, "N": DATUM})
    n_bu = len(buchungen)
    s = n_bu + 2
    storno = Buchungsstatus.STORNIERT.label
    bezahlt = Buchungsstatus.BEZAHLT.label
    bu[f"A{s}"] = "Summe (ohne stornierte)"
    bu[f"D{s}"] = f'=SUMIF({_bereich("O", n_bu)},"<>{storno}",{_bereich("D", n_bu)})'
    bu[f"J{s}"] = f'=SUMIF({_bereich("O", n_bu)},"<>{storno}",{_bereich("J", n_bu)})'
    bu[f"M{s}"] = (f'=SUMIF({_bereich("O", n_bu)},"<>{storno}",{_bereich("M", n_bu)})'
                   f'-SUMIF({_bereich("O", n_bu)},"{bezahlt}",{_bereich("M", n_bu)})')
    bu[f"J{s}"].number_format = bu[f"M{s}"].number_format = EURO
    for sp in "ABCDEFGHIJKLMNOPQR":
        bu[f"{sp}{s}"].font, bu[f"{sp}{s}"].border = FETT, LINIE_OBEN

    # Weitere Kosten
    ko = wb.create_sheet("Kosten")
    _kopfzeile(ko, ["Art", "Betrag €", "Notiz"], [14, 12, 50])
    kosten = list(reise.kosten.all())
    for k in kosten:
        ko.append([k.get_art_display(), k.betrag, k.notiz])
    _formate(ko, len(kosten), {"B": EURO})

    # Links
    li = wb.create_sheet("Links")
    _kopfzeile(li, ["Titel", "URL"], [40, 90])
    for i, link in enumerate(reise.links.all(), start=2):
        li.append([link.titel, link.url])
        li[f"B{i}"].hyperlink = link.url
        li[f"B{i}"].style = "Hyperlink"

    # Übersicht – alles per Formel aus den anderen Blättern
    ue.column_dimensions["A"].width = 28
    ue.column_dimensions["B"].width = 22
    ue.column_dimensions["C"].width = 40
    ue["A1"] = reise.titel
    ue["A1"].font = Font(bold=True, size=16, color=BLAU)
    ue["A2"] = f"Urlaub {reise.urlaub_beginn:%d.%m.%Y} – {reise.urlaub_ende:%d.%m.%Y}"
    if reise.hat_eigenen_urlaub:
        ue["A2"] = ue["A2"].value + f", unterwegs {reise.von:%d.%m.} – {reise.bis:%d.%m.%Y}" + (
            " (Abfahrt nach Feierabend)" if reise.abfahrt_feierabend else "")
    ue["C1"] = f"Stand: {stand:%d.%m.%Y, %H:%M} Uhr"
    ue["C1"].font = Font(italic=True, color="666666")
    etappen_summe = n_et + 2
    buchung_summe = n_bu + 2
    zeilen = [
        ("Urlaubstage", reise.urlaubstage or 0, "0", f"Mo–Fr ohne Feiertage {reise.bundesland}, berechnet"),
        ("Fahrzeug", reise.fahrzeugprofil.name if reise.fahrzeugprofil else "–", None, ""),
        ("Kraftstoffpreis (€/L)", reise.dieselpreis, '0.000 "€"', "Eingabe – ändert die Kraftstoffkosten"),
        ("Entfernung", f"=Etappen!I{etappen_summe}", KM, ""),
        ("Fahrzeit", f"=Etappen!F{etappen_summe}", DAUER, "ohne Pausen"),
        ("Kraftstoff (Liter)", f"=Etappen!K{etappen_summe}", "0.0", "je Etappe km × L/100 km bzw. Liter laut TomTom-Verbrauchskurve"),
        ("Kraftstoff", "=ROUND(B9*B6,2)", EURO, ""),
        ("Übernachtungen", f"=Buchungen!J{buchung_summe}", EURO, ""),
    ]
    for art, label in Kostenart.choices:
        if any(k.art == art for k in kosten):
            zeilen.append((label, f'=SUMIF(Kosten!A:A,"{label}",Kosten!B:B)', EURO, ""))
    for z, (name, wert, fmt, notiz) in enumerate(zeilen, start=4):
        ue[f"A{z}"], ue[f"B{z}"], ue[f"C{z}"] = name, wert, notiz
        if fmt:
            ue[f"B{z}"].number_format = fmt
        ue[f"C{z}"].font = Font(italic=True, color="666666")
    letzte = 3 + len(zeilen)
    g = letzte + 1
    ue[f"A{g}"], ue[f"B{g}"] = "Gesamtkosten", f"=SUM(B10:B{letzte})"
    ue[f"A{g + 1}"], ue[f"B{g + 1}"] = "davon noch offen", f"=Buchungen!M{buchung_summe}"
    ue[f"A{g + 2}"], ue[f"B{g + 2}"] = "Nächte", f"=Buchungen!D{buchung_summe}"
    for z in (g, g + 1):
        ue[f"B{z}"].number_format = EURO
    ue[f"A{g}"].font = ue[f"B{g}"].font = Font(bold=True, size=12)
    ue[f"A{g}"].border = ue[f"B{g}"].border = LINIE_OBEN

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# --- PDF ---------------------------------------------------------------------


def pdf(reise, kontext, stand):
    from weasyprint import HTML

    html = render_to_string("reisen/pdf.html", {
        **kontext, "schriften": (settings.BASE_DIR / "static" / "fonts").as_uri(), "stand": stand,
        "zeitraum": zeitraum_text(reise.von, reise.bis)})
    return HTML(string=html, base_url=str(settings.BASE_DIR)).write_pdf()
