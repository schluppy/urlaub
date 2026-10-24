from datetime import date, time, timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse

from konten.models import Person
from orte.models import Ort
from reisen.formulare import DauerFeld
from reisen.models import Etappe, Reise

from .hilfen import fam


class DauerFeldTest(TestCase):
    def test_stunden_minuten(self):
        f = DauerFeld()
        self.assertEqual(f.clean("1:15"), timedelta(hours=1, minutes=15))   # nicht 1 min 15 s
        self.assertEqual(f.clean("09:27"), timedelta(hours=9, minutes=27))
        self.assertEqual(f.clean("01:00:00"), timedelta(hours=1))
        self.assertEqual(f.prepare_value(timedelta(hours=8, minutes=37, seconds=40)), "08:38")
        for falsch in ["abc", "1:75", "115"]:
            with self.subTest(falsch), self.assertRaises(Exception):
                f.clean(falsch)

    def test_unveraendert(self):
        self.assertFalse(DauerFeld().has_changed(timedelta(), "00:00"))
        self.assertTrue(DauerFeld().has_changed(timedelta(), "00:30"))


def leere_inlines(**ausser):
    daten = {}
    for prefix in ["etappen", "buchungen", "kosten", "links", "erinnerungen", "tracks"]:
        daten.update({f"{prefix}-TOTAL_FORMS": "0", f"{prefix}-INITIAL_FORMS": "0",
                      f"{prefix}-MIN_NUM_FORMS": "0", f"{prefix}-MAX_NUM_FORMS": "1000"})
    daten.update(ausser)
    return daten


class AdminEingabeTest(TestCase):
    def setUp(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        self.zuhause = Ort.objects.create(familie=fam(), name="Zuhause", typ="zuhause")
        self.rhein = Ort.objects.create(familie=fam(), name="Rheincamping")
        self.see = Ort.objects.create(familie=fam(), name="Sommerreise")

    def test_formular_mit_pickern(self):
        r = Reise.objects.create(familie=fam(), titel="Sommerreise", von=date(2027, 8, 27), bis=date(2027, 9, 20))
        Etappe.objects.create(reise=r, datum=date(2027, 8, 27), von_ort=self.zuhause, nach_ort=self.rhein,
                              startzeit=time(6), pause=timedelta(hours=1))
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[r.pk])).content.decode()
        self.assertIn('type="date" name="von" value="2027-08-27"', seite)
        self.assertIn('type="time" name="etappen-0-startzeit" value="06:00"', seite)
        self.assertIn('type="time" name="etappen-0-pause" value="01:00"', seite)
        self.assertIn("reisen/eingabe.js", seite)
        self.assertIn("Ganze Reise löschen", seite)                       # roter Knopf eindeutig
        self.assertIn('class="loesch-hinweis"', seite)
        self.assertIn('value="Sichern" class="default" name="_continue"', seite)
        self.assertIn('value="Sichern und zurück zum Urlaubsplaner" name="_planer"', seite)
        self.assertNotIn('name="_addanother"', seite)
        self.assertNotIn('name="_save"', seite)
        self.assertNotIn('name="etappen-0-fahrzeugprofil"', seite)           # Spalte entfernt
        self.assertNotIn("delete-related", seite)                           # kein ✕ neben Auswahlfeldern
        self.assertIn("change-related", seite)                              # Stift bleibt

    def test_etappen_nach_datum_sortiert(self):
        daten = leere_inlines(**{
            "titel": "Sommerreise", "von": "2027-08-27", "bis": "2027-09-20", "urlaub_von": "", "urlaub_bis": "",
            "fahrzeugprofil": "", "verbrauch_l_100km": "", "dieselpreis": "1.80", "vignetten_meiden": "",
            "im_kalender": "on", "notizen": "",
            "etappen-TOTAL_FORMS": "2",
            # bewusst in falscher Reihenfolge eingegeben
            "etappen-0-datum": "2027-08-28", "etappen-0-von_ort": self.rhein.pk, "etappen-0-nach_ort": self.see.pk,
            "etappen-0-startzeit": "07:00", "etappen-0-pause": "0:45", "etappen-0-fahrzeit": "",
            "etappen-1-datum": "2027-08-27", "etappen-1-von_ort": self.zuhause.pk, "etappen-1-nach_ort": self.rhein.pk,
            "etappen-1-startzeit": "06:00", "etappen-1-pause": "1:15", "etappen-1-fahrzeit": "9:25",
        })
        antwort = self.client.post(reverse("admin:reisen_reise_add"), daten)
        self.assertEqual(antwort.status_code, 302, antwort.content.decode()[:3000])
        etappen = list(Reise.objects.get().etappen.all())
        self.assertEqual([e.datum.day for e in etappen], [27, 28])
        self.assertEqual(etappen[0].pause, timedelta(hours=1, minutes=15))
        self.assertEqual(etappen[0].fahrzeit, timedelta(hours=9, minutes=25))
        self.assertEqual(etappen[0].ankunft.strftime("%H:%M"), "16:40")

    def test_sichern_bleibt_oder_fuehrt_zum_urlaubsplaner(self):
        r = Reise.objects.create(familie=fam(), titel="Sommerreise", von=date(2027, 8, 27), bis=date(2027, 9, 20))
        daten = leere_inlines(titel="Sommerreise", von="2027-08-27", bis="2027-09-20", urlaub_von="", urlaub_bis="",
                              fahrzeugprofil="", verbrauch_l_100km="", dieselpreis="1.80", vignetten_meiden="",
                              im_kalender="on", notizen="Fähre")
        url = reverse("admin:reisen_reise_change", args=[r.pk])
        antwort = self.client.post(url, {**daten, "_continue": "Sichern"})
        self.assertRedirects(antwort, url, fetch_redirect_response=False)
        antwort = self.client.post(url, {**daten, "_planer": "1"})
        self.assertRedirects(antwort, reverse("reisen:reise", args=[r.pk]), fetch_redirect_response=False)
        antwort = self.client.post(reverse("admin:reisen_reise_add"), {**daten, "titel": "Neu", "_planer": "1"})
        self.assertRedirects(antwort, reverse("reisen:reise", args=[Reise.objects.get(titel="Neu").pk]),
                             fetch_redirect_response=False)

    def test_gleicher_tag_nach_startzeit(self):
        r = Reise.objects.create(familie=fam(), titel="X", von=date(2027, 8, 27), bis=date(2027, 9, 20))
        Etappe.objects.create(reise=r, datum=date(2027, 9, 19), startzeit=time(7), von_ort=self.see,
                              nach_ort=self.zuhause)
        Etappe.objects.create(reise=r, datum=date(2027, 8, 27), startzeit=time(14), von_ort=self.rhein,
                              nach_ort=self.see)
        Etappe.objects.create(reise=r, datum=date(2027, 8, 27), startzeit=time(6), von_ort=self.zuhause,
                              nach_ort=self.rhein)
        self.assertEqual([(e.datum.day, e.startzeit.hour) for e in r.etappen.all()], [(27, 6), (27, 14), (19, 7)])
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[r.pk])).content.decode()
        self.assertNotIn('name="etappen-0-reihenfolge"', seite)
        self.assertLess(seite.index('value="2027-08-27"'), seite.index('value="2027-09-19"'))


class BuchungsauswahlTest(TestCase):
    """Buchungen: nur Übernachtungsorte aus den Etappen, Daten aus den Aufenthalten."""

    def setUp(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        self.zuhause = Ort.objects.create(familie=fam(), name="Zuhause", typ="zuhause")
        self.rhein = Ort.objects.create(familie=fam(), name="Rheincamping")
        self.see = Ort.objects.create(familie=fam(), name="Seecamping")
        self.fremd = Ort.objects.create(familie=fam(), name="Anderer Platz")
        self.reise = Reise.objects.create(familie=fam(), titel="Sommerreise", von=date(2027, 8, 27), bis=date(2027, 9, 19))
        for tag, von, nach in [(date(2027, 8, 27), self.zuhause, self.rhein),
                               (date(2027, 8, 28), self.rhein, self.see),
                               (date(2027, 9, 18), self.see, self.rhein),
                               (date(2027, 9, 19), self.rhein, self.zuhause)]:
            Etappe.objects.create(reise=self.reise, datum=tag, von_ort=von, nach_ort=nach)

    def test_aufenthalte(self):
        self.assertEqual([(o.name, a.day, b.day) for o, a, b in self.reise.aufenthalte()],
                         [("Rheincamping", 27, 28), ("Seecamping", 28, 18), ("Rheincamping", 18, 19)])

    def test_auswahl_und_vorbelegung(self):
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[self.reise.pk])).content.decode()
        auswahl = seite[seite.index('name="buchungen-__prefix__-ort"'):]
        auswahl = auswahl[:auswahl.index("</select>")]
        self.assertIn("Rheincamping", auswahl)
        self.assertIn("Seecamping", auswahl)
        self.assertNotIn("Anderer Platz", auswahl)
        self.assertNotIn("Zuhause", auswahl)
        self.assertIn('<script id="aufenthalte" type="application/json">', seite)
        self.assertIn(f'{{"ort": {self.rhein.pk}, "von": "2027-09-18", "bis": "2027-09-19"}}', seite)

    def test_bereits_gebuchter_ort_bleibt_waehlbar(self):
        from reisen.models import Buchung
        Buchung.objects.create(reise=self.reise, ort=self.fremd, von=date(2027, 8, 27), bis=date(2027, 8, 28))
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[self.reise.pk])).content.decode()
        self.assertIn("Anderer Platz", seite[seite.index('name="buchungen-0-ort"'):][:2000])

    def test_neue_reise_ohne_etappen(self):
        seite = self.client.get(reverse("admin:reisen_reise_add")).content.decode()
        self.assertIn("Neue Etappen erst speichern", seite)


class LinkTest(TestCase):
    def test_links_ohne_nummer_in_eingabereihenfolge(self):
        from reisen.models import Link
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        r = Reise.objects.create(familie=fam(), titel="Schottland", von=date(2026, 5, 29), bis=date(2026, 6, 21))
        for titel in ["Fähre DFDS", "ETA Großbritannien", "Blog"]:
            Link.objects.create(reise=r, titel=titel, url="https://example.org")
        self.assertEqual([link.titel for link in r.links.all()], ["Fähre DFDS", "ETA Großbritannien", "Blog"])
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[r.pk])).content.decode()
        self.assertNotIn('name="links-0-reihenfolge"', seite)
        self.assertIn('name="links-0-titel"', seite)


class BuchungOhnePlatzTest(TestCase):
    """Neue Etappen und Buchung in einem Schritt: die Buchungszeile ohne Platz blockiert das Speichern nicht."""

    def test_etappen_werden_gespeichert(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        zuhause = Ort.objects.create(familie=fam(), name="Zuhause", typ="zuhause")
        see = Ort.objects.create(familie=fam(), name="Seecamping")
        daten = leere_inlines(**{
            "titel": "Rundreise", "von": "2027-08-01", "bis": "2027-08-20", "urlaub_von": "", "urlaub_bis": "",
            "fahrzeugprofil": "", "verbrauch_l_100km": "", "dieselpreis": "1.80", "vignetten_meiden": "",
            "im_kalender": "on", "notizen": "",
            "etappen-TOTAL_FORMS": "2",
            "etappen-0-datum": "2027-08-01", "etappen-0-von_ort": zuhause.pk, "etappen-0-nach_ort": see.pk,
            "etappen-0-pause": "0:00",
            "etappen-1-datum": "2027-08-20", "etappen-1-von_ort": see.pk, "etappen-1-nach_ort": zuhause.pk,
            "etappen-1-pause": "0:00",
            # vorbelegt, aber Platz noch nicht wählbar
            "buchungen-TOTAL_FORMS": "1", "buchungen-0-ort": "", "buchungen-0-von": "2027-08-01",
            "buchungen-0-bis": "", "buchungen-0-waehrung": "EUR", "buchungen-0-wechselkurs": "1",
            "buchungen-0-anzahlung": "0", "buchungen-0-status": "geplant", "buchungen-0-preis_pro_nacht": "30",
        })
        antwort = self.client.post(reverse("admin:reisen_reise_add"), daten)
        self.assertEqual(antwort.status_code, 302, antwort.content.decode()[:3000])
        r = Reise.objects.get()
        self.assertEqual(r.etappen.count(), 2)
        self.assertEqual(r.buchungen.count(), 0)
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[r.pk])).content.decode()
        self.assertIn("Kraftstoffpreis", seite)
        self.assertIn("Kraftstoff", seite[seite.index("Berechnet"):])


class SucheZuruecksetzenTest(TestCase):
    def test_knopf_nur_mit_suchbegriff_und_filter_bleibt(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        Ort.objects.create(familie=fam(), name="Seecamping")
        Ort.objects.create(familie=fam(), name="Bergcamping")
        url = reverse("admin:orte_ort_changelist")
        self.assertNotContains(self.client.get(url), "Zurücksetzen")
        seite = self.client.get(url, {"q": "See", "typ__exact": "camping"})
        self.assertContains(seite, "Zurücksetzen")
        self.assertContains(seite, 'href="?typ__exact=camping" class="button suche-zuruecksetzen"')
        self.assertNotContains(seite, "Bergcamping</a>")


class PopupTest(TestCase):
    def test_skript_in_allen_masken(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        for url in [reverse("admin:orte_ort_add") + "?_to_field=id&_popup=1", reverse("admin:reisen_reise_add")]:
            with self.subTest(url):
                self.assertContains(self.client.get(url), "verwaltung/popup.js")


class FremdwaehrungTest(TestCase):
    """Buchungen als Blöcke; Währung und Kurs zählen nur mit Haken „Fremdwährung“; Stellplatz."""

    def setUp(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        zuhause = Ort.objects.create(familie=fam(), name="Zuhause", typ="zuhause")
        self.see = Ort.objects.create(familie=fam(), name="Seecamping")
        self.r = Reise.objects.create(familie=fam(), titel="Sommer", von=date(2027, 8, 1), bis=date(2027, 8, 5))
        Etappe.objects.create(reise=self.r, datum=date(2027, 8, 1), von_ort=zuhause, nach_ort=self.see)
        Etappe.objects.create(reise=self.r, datum=date(2027, 8, 5), von_ort=self.see, nach_ort=zuhause)

    def speichern(self, **buchung):
        e = list(self.r.etappen.all())
        daten = leere_inlines(**{
            "titel": "Sommer", "von": "2027-08-01", "bis": "2027-08-05", "urlaub_von": "", "urlaub_bis": "",
            "fahrzeugprofil": "", "verbrauch_l_100km": "", "dieselpreis": "1.80", "vignetten_meiden": "",
            "im_kalender": "on", "notizen": "",
            "etappen-TOTAL_FORMS": "2", "etappen-INITIAL_FORMS": "2",
            **{f"etappen-{i}-{k}": v for i, x in enumerate(e) for k, v in
               {"id": x.pk, "reise": self.r.pk, "datum": x.datum.isoformat(), "von_ort": x.von_ort_id,
                "nach_ort": x.nach_ort_id, "pause": "0:00"}.items()},
            "buchungen-TOTAL_FORMS": "1", "buchungen-0-ort": self.see.pk, "buchungen-0-von": "2027-08-01",
            "buchungen-0-bis": "2027-08-05", "buchungen-0-preis_pro_nacht": "30", "buchungen-0-anzahlung": "0",
            "buchungen-0-status": "geplant", **{f"buchungen-0-{k}": v for k, v in buchung.items()},
        })
        return self.client.post(reverse("admin:reisen_reise_change", args=[self.r.pk]), daten)

    def test_ohne_haken_euro(self):
        self.assertEqual(self.speichern(waehrung="GBP", wechselkurs="0.85").status_code, 302)
        b = self.r.buchungen.get()
        self.assertEqual((b.waehrung, b.wechselkurs), ("EUR", 1))

    def test_mit_haken_fremdwaehrung(self):
        self.assertEqual(self.speichern(fremdwaehrung="on", waehrung="GBP", wechselkurs="0.85").status_code, 302)
        b = self.r.buchungen.get()
        self.assertEqual((b.waehrung, str(b.wechselkurs)), ("GBP", "0.8500"))
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[self.r.pk])).content.decode()
        self.assertIn('class="inline-related', seite[seite.index('id="buchungen-group"'):])   # Blöcke
        self.assertRegex(seite, r'<input type="checkbox" name="buchungen-0-fremdwaehrung"[^>]*checked')

    def test_preis_als_rechnung(self):
        for feld, text, soll in [("preis_pro_nacht", "=24+(2*8)+4+9", "53.00"), ("gesamtpreis", "=4*(30,5)", "122.00"),
                                 ("preis_pro_nacht", "28,40", "28.40")]:
            with self.subTest(text):
                self.r.buchungen.all().delete()
                preise = {"preis_pro_nacht": "", "gesamtpreis": "", feld: text}
                self.assertEqual(self.speichern(waehrung="EUR", wechselkurs="1", **preise).status_code, 302)
                self.assertEqual(str(getattr(self.r.buchungen.get(), feld)), soll)
        seite = self.client.get(reverse("admin:reisen_reise_change", args=[self.r.pk])).content.decode()
        self.assertRegex(seite, r'name="buchungen-0-preis_pro_nacht" value="28,40"')   # Dezimalkomma

    def test_unverstaendliche_rechnung_ist_fehler(self):
        for text in ["=24+", "=2**8", "=__import__('os')", "=1/0"]:
            with self.subTest(text):
                self.assertContains(self.speichern(preis_pro_nacht=text, waehrung="EUR", wechselkurs="1"), "Rechnung nicht verständlich")

    def test_haken_ohne_waehrung_ist_fehler(self):
        antwort = self.speichern(fremdwaehrung="on", waehrung="EUR", wechselkurs="1")
        self.assertEqual(antwort.status_code, 200)
        self.assertContains(antwort, "Bitte die Fremdwährung angeben")

    def test_stellplatz_eingabe_seite_excel(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from reisen.dokumente import excel, stand_jetzt

        antwort = self.speichern(buchungsnummer="A-1", stellplatz="47", waehrung="EUR", wechselkurs="1")
        self.assertEqual(antwort.status_code, 302)
        self.assertEqual(self.r.buchungen.get().stellplatz, "47")
        self.assertContains(self.client.get(reverse("reisen:reise", args=[self.r.pk])), "Nr. A-1 · Stellplatz 47")
        ws = load_workbook(BytesIO(excel(Reise.objects.get(pk=self.r.pk), stand_jetzt())))["Buchungen"]
        self.assertEqual((ws["P1"].value, ws["Q1"].value, ws["Q2"].value, ws["R1"].value),
                         ("Buchungsnr.", "Stellplatz", "47", "Website"))
