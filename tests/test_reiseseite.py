from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from fahrzeuge.models import Fahrzeugprofil
from konten.models import Person
from orte.models import Ort
from reisen.models import Buchung, Etappe, Kostenposition, Link, Reise

from .hilfen import fam


def schottland():
    """Ausschnitt aus dem Blatt „Schottland“ der Tabelle 2026."""
    profil = Fahrzeugprofil.objects.create(familie=fam(), name="Gespann", verbrauch_l_100km=Decimal("12.2"))
    zuhause = Ort.objects.create(familie=fam(), name="Zuhause", typ="zuhause", breite=Decimal("54.32"), laenge=Decimal("10.13"))
    ams = Ort.objects.create(familie=fam(), name="DFDS Amsterdam Terminal", typ="faehre", breite=Decimal("52.46"), laenge=Decimal("4.6"))
    drum = Ort.objects.create(familie=fam(), name="Küstencamping", website="https://kuestencamping.example")
    gran = Ort.objects.create(familie=fam(), name="Highland Caravan Park")
    r = Reise.objects.create(familie=fam(), titel="Schottland", von=date(2026, 5, 29), bis=date(2026, 6, 21),
                             fahrzeugprofil=profil, dieselpreis=Decimal("2.06"), notizen="Fähre buchbar ab Oktober")
    Etappe.objects.create(reise=r, datum=date(2026, 5, 29), von_ort=zuhause, nach_ort=ams,
                          startzeit=time(5, 30), fahrzeit=timedelta(hours=6, minutes=34),
                          pause=timedelta(hours=1, minutes=15), entfernung_km=Decimal(517))
    Etappe.objects.create(reise=r, datum=date(2026, 6, 2), von_ort=drum, nach_ort=gran,
                          startzeit=time(10), fahrzeit=timedelta(hours=4, minutes=12))
    Buchung.objects.create(reise=r, ort=drum, von=date(2026, 5, 30), bis=date(2026, 6, 2),
                           preis_pro_nacht=Decimal(48), waehrung="GBP", wechselkurs=Decimal("0.8653"),
                           anzahlung=Decimal(36), status="angezahlt", rest_faellig_am=date(2026, 5, 1))
    Buchung.objects.create(reise=r, ort=gran, von=date(2026, 6, 3), bis=date(2026, 6, 4),
                           preis_pro_nacht=Decimal("28.40"), waehrung="GBP", wechselkurs=Decimal("0.8653"))
    Kostenposition.objects.create(reise=r, art="faehre", betrag=Decimal(1442), notiz="DFDS hin und zurück")
    Link.objects.create(reise=r, titel="ETA Großbritannien",
                        url="https://www.gov.uk/guidance/apply-for-an-electronic-travel-authorisation-eta")
    return r


class HinweiseTest(TestCase):
    def test_hinweise(self):
        r = schottland()
        texte = " | ".join(t for _, t in r.hinweise(heute=date(2026, 5, 10)))
        self.assertIn("1 Etappe ohne Entfernung", texte)
        self.assertIn("Nächte ohne Buchung: 29.05., 02.06., 04.06.", texte)
        self.assertIn("Restzahlung Küstencamping war am 01.05.2026 fällig", texte)

    def test_kostenaufstellung(self):
        r = schottland()
        self.assertEqual([n for n, _ in r.kostenaufstellung()], ["Kraftstoff", "Übernachtungen", "Fähre"])
        self.assertEqual(r.etappen.first().karten_link,
                         "https://www.google.com/maps/dir/?api=1&origin=54.320000,10.130000"
                         "&destination=52.460000,4.600000&travelmode=driving")


class SeiteTest(TestCase):
    def setUp(self):
        self.client.force_login(Person.objects.create_superuser("t@example.org", "T", familie=fam()))
        self.reise = schottland()

    def test_reiseseite(self):
        antwort = self.client.get(reverse("reisen:reise", args=[self.reise.pk]))
        self.assertContains(antwort, "Küstencamping")
        self.assertContains(antwort, "144,00 GBP")       # 3 × 48 £
        self.assertContains(antwort, "166,42 €")          # 144 £ / 0,8653
        self.assertContains(antwort, "13:19")             # 5:30 + 6:34 + 1:15
        self.assertContains(antwort, "1.442,00 €")
        self.assertContains(antwort, "ETA Großbritannien")
        self.assertContains(antwort, "<svg class=\"kostenbalken\"")
        uebersicht = self.client.get(reverse("portal:jahr", args=[2026]))
        self.assertContains(uebersicht, reverse("reisen:reise", args=[self.reise.pk]))

    def test_unbekannte_reise(self):
        self.assertEqual(self.client.get(reverse("reisen:reise", args=[999])).status_code, 404)

    @override_settings(TOMTOM_API_KEY="test")
    def test_route_knopf(self):
        with mock.patch("reisen.views.reise_routen", return_value=(2, [])) as routen:
            antwort = self.client.post(reverse("reisen:route", args=[self.reise.pk]), follow=True)
        routen.assert_called_once()
        self.assertContains(antwort, "2 Etappen mit TomTom berechnet")
        self.assertEqual(self.client.get(reverse("reisen:route", args=[self.reise.pk])).status_code, 405)

    def test_pdf(self):
        antwort = self.client.get(reverse("reisen:pdf", args=[self.reise.pk]))
        self.assertEqual(antwort["Content-Type"], "application/pdf")
        self.assertTrue(antwort.content.startswith(b"%PDF"))
        self.assertRegex(antwort["Content-Disposition"], r'filename="2026-05-29_Schottland_Stand_\d{4}-\d{2}-\d{2}_\d{4}\.pdf"')

    def test_excel_mit_formeln(self):
        antwort = self.client.get(reverse("reisen:excel", args=[self.reise.pk]))
        wb = load_workbook(BytesIO(antwort.content))
        self.assertEqual(wb.sheetnames, ["Übersicht", "Etappen", "Buchungen", "Kosten", "Links"])
        et, bu, ue = wb["Etappen"], wb["Buchungen"], wb["Übersicht"]
        self.assertEqual(et["H2"].value, '=IF(OR(E2="",F2=""),"",E2+F2+G2)')
        self.assertEqual(et["K2"].value, '=IF(OR(I2="",J2=""),0,I2*J2/100)')
        self.assertEqual(bu["D2"].value, "=C2-B2")
        self.assertEqual(bu["J2"].value, "=ROUND(G2/I2,2)")
        self.assertEqual(ue["A10"].value, "Kraftstoff")
        self.assertEqual(ue["B10"].value, "=ROUND(B9*B6,2)")
        self.assertEqual(ue["A12"].value, "Fähre")
        self.assertEqual(ue["A13"].value, "Gesamtkosten")
        self.assertEqual(ue["B13"].value, "=SUM(B10:B12)")


class PreisTest(TestCase):
    def setUp(self):
        self.reise = Reise.objects.create(familie=fam(), titel="Sommerreise", von=date(2027, 8, 27), bis=date(2027, 9, 19))
        self.ort = Ort.objects.create(familie=fam(), name="Camping Seecamping")

    def buchung(self, **preis):
        return Buchung(reise=self.reise, ort=self.ort, von=date(2027, 8, 28), bis=date(2027, 9, 18), **preis)

    def test_gesamtpreis(self):
        b = self.buchung(gesamtpreis=Decimal(563), anzahlung=Decimal(100))
        b.full_clean()
        self.assertEqual((b.naechte, b.summe, b.rest, b.preis_je_nacht), (21, 563, 463, Decimal("26.81")))
        self.assertTrue(b.ist_pauschal)

    def test_preis_pro_nacht(self):
        b = self.buchung(preis_pro_nacht=Decimal("26.50"))
        self.assertEqual((b.summe, b.preis_je_nacht, b.ist_pauschal), (Decimal("556.50"), Decimal("26.50"), False))

    def test_ohne_preis(self):
        self.assertEqual(self.buchung().summe, 0)

    def test_beide_passend_erlaubt_widerspruch_nicht(self):
        self.buchung(preis_pro_nacht=Decimal("26.50"), gesamtpreis=Decimal("556.50")).full_clean()
        with self.assertRaisesMessage(ValidationError, "entweder Preis/Nacht oder Gesamtpreis"):
            self.buchung(preis_pro_nacht=Decimal(30), gesamtpreis=Decimal(563)).full_clean()

    def test_anzeige_und_excel(self):
        self.buchung(gesamtpreis=Decimal(563)).save()
        Person.objects.create_superuser("t@example.org", "T", familie=fam())
        self.client.force_login(Person.objects.get())
        seite = self.client.get(reverse("reisen:reise", args=[self.reise.pk]))
        self.assertContains(seite, "pauschal")
        self.assertContains(seite, "≈ 26,81 €/Nacht")
        self.assertContains(seite, "563,00 €")
        wb = load_workbook(BytesIO(self.client.get(reverse("reisen:excel", args=[self.reise.pk])).content))
        bu = wb["Buchungen"]
        self.assertEqual((bu["E1"].value, bu["F1"].value), ("Preis/Nacht", "Gesamtpreis"))
        self.assertEqual((bu["E2"].value, bu["F2"].value, bu["G2"].value), (None, 563, '=IF(F2<>"",F2,D2*E2)'))


class EtappenLinksTest(TestCase):
    def test_plaetze_in_etappen_verlinkt(self):
        f = fam()
        zuhause = Ort.objects.create(familie=f, name="Daheim", typ="zuhause", website="https://daheim.example/")
        see = Ort.objects.create(familie=f, name="Seecamping", website="https://see.example/")
        r = Reise.objects.create(familie=f, titel="Kurz", von=date(2027, 8, 1), bis=date(2027, 8, 3))
        Etappe.objects.create(reise=r, datum=date(2027, 8, 1), von_ort=zuhause, nach_ort=see)
        self.client.force_login(Person.objects.create_user("a@example.org", "A", familie=f, is_staff=True))
        seite = self.client.get(reverse("reisen:reise", args=[r.pk]))
        self.assertContains(seite, '<a href="https://see.example/" target="_blank" rel="noopener">Seecamping</a>')
        self.assertNotContains(seite, "daheim.example")                     # Zuhause nie verlinkt


class AnleitungTest(TestCase):
    def test_ohne_anmeldung_lesbar_und_im_menue(self):
        seite = self.client.get(reverse("portal:anleitung"))
        self.assertContains(seite, "Anmelden ohne Passwort")
        self.assertContains(seite, 'aria-current="page">Anleitung</a>')
        self.assertContains(seite, "15 Minuten")
        self.assertNotContains(seite, ">Bearbeiten</a>")                   # ohne Anmeldung keine weiteren Menüpunkte
        self.assertContains(self.client.get(reverse("portal:anmelden")), reverse("portal:anleitung"))

    def test_menue_angemeldet_und_anker(self):
        import re

        self.client.force_login(Person.objects.create_user("a@example.org", "A", familie=fam(), is_staff=True))
        seite = self.client.get(reverse("portal:start")).content.decode()
        self.assertIn(f'href="{reverse("portal:anleitung")}"', seite)
        self.assertIn('aria-current="page">Reisen</a>', seite)
        anleitung = self.client.get(reverse("portal:anleitung")).content.decode()
        self.assertNotIn('aria-current="page">Reisen</a>', anleitung)
        anker = set(re.findall(r'href="#([\w-]+)"', anleitung))
        ziele = set(re.findall(r'id="([\w-]+)"', anleitung))
        self.assertTrue(anker and anker <= ziele, anker - ziele)              # Inhaltsverzeichnis ohne tote Links


class AdacLinkTest(TestCase):
    def test_link_je_etappe_mit_fahrzeugart(self):
        f = fam()
        a = Ort.objects.create(familie=f, name="Start", breite=Decimal("54.1"), laenge=Decimal("10.25"))
        b = Ort.objects.create(familie=f, name="Ziel", breite=Decimal("45.123456"), laenge=Decimal("7.5"))
        ohne = Ort.objects.create(familie=f, name="Ohne Koordinaten")
        mobil = Fahrzeugprofil.objects.create(familie=f, name="Mobil", art="motorhome")
        r = Reise.objects.create(familie=f, titel="Kurz", von=date(2027, 8, 1), bis=date(2027, 8, 3))
        e = Etappe.objects.create(reise=r, datum=date(2027, 8, 1), von_ort=a, nach_ort=b)
        self.assertEqual(e.adac_link, "https://maps.adac.de/route?places=54.10000_10.25000_2_0,"
                                      "45.12346_7.50000_6_0&vehicle-type=trailer")    # ohne Profil: Gespann
        r.fahrzeugprofil = mobil
        r.save()
        e = Etappe.objects.get(pk=e.pk)
        self.assertTrue(e.adac_link.endswith("&vehicle-type=motorhome"))
        self.assertEqual(Etappe(reise=r, datum=r.von, von_ort=a, nach_ort=ohne).adac_link, "")
        self.client.force_login(Person.objects.create_user("a@example.org", "A", familie=f, is_staff=True))
        seite = self.client.get(reverse("reisen:reise", args=[r.pk]))
        self.assertContains(seite, 'title="Route mit Maut in ADAC Maps">ADAC</a>')
