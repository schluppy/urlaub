from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from fahrzeuge.models import Fahrzeugprofil
from konten.models import Person
from orte.models import Ort
from reisen.dokumente import excel, stand_jetzt
from reisen.models import Etappe, Reise

from .hilfen import fam


class VerbrauchskurveTest(TestCase):
    """Liter laut TomTom (Verbrauchskurve) gehen vor; ohne sie zählt der Ø-Verbrauch."""

    def setUp(self):
        self.f = fam()
        self.profil = Fahrzeugprofil.objects.create(familie=self.f, name="Gespann", verbrauch_l_100km=Decimal("11.5"),
                                                    verbrauchskurve="80,10:90,11:100,12")
        a, b, c = (Ort.objects.create(familie=self.f, name=n) for n in ("Zuhause", "Seecamping", "Bergcamping"))
        self.r = Reise.objects.create(familie=self.f, titel="Sommer", von=date(2027, 8, 1), bis=date(2027, 8, 20),
                                      fahrzeugprofil=self.profil, dieselpreis=Decimal("2.00"))
        self.mit = Etappe.objects.create(reise=self.r, datum=date(2027, 8, 1), von_ort=a, nach_ort=b,
                                         entfernung_km=Decimal("600"), routing_details={"verbrauch_l": "65.4"})
        self.ohne = Etappe.objects.create(reise=self.r, datum=date(2027, 8, 20), von_ort=b, nach_ort=c,
                                          entfernung_km=Decimal("200"))

    def reise(self):
        return Reise.objects.get(pk=self.r.pk)

    def test_tomtom_liter_und_rest_mit_durchschnitt(self):
        self.assertEqual(self.mit.dieselliter, Decimal("65.4"))
        self.assertEqual(self.ohne.dieselliter, Decimal("23"))              # 200 km × 11,5 / 100
        r = self.reise()
        self.assertEqual(r.dieselliter, Decimal("88.4"))
        self.assertEqual(r.dieselkosten, Decimal("176.80"))
        self.assertEqual(r.diesel_quelle, "Ø 11,1 L/100 km, 1 von 2 Etappen mit Verbrauchskurve TomTom")

    def test_alle_etappen_mit_tomtom(self):
        self.ohne.routing_details = {"verbrauch_l": "21.0"}
        self.ohne.save()
        self.assertEqual(self.reise().diesel_quelle, "Ø 10,8 L/100 km, Verbrauchskurve TomTom")

    def test_ohne_kurve_im_profil_nur_durchschnitt(self):
        self.profil.verbrauchskurve = ""
        self.profil.save()
        r = self.reise()
        self.assertIsNone(r.etappen.first().tomtom_liter)
        self.assertEqual(r.dieselliter, Decimal("92"))                      # 800 km × 11,5 / 100
        self.assertEqual(r.diesel_quelle, "11,5 L/100 km")

    def test_reiseseite_und_excel(self):
        person = Person.objects.create_user("anna@example.org", "Anna", familie=self.f, is_staff=True)
        self.client.force_login(person)
        seite = self.client.get(reverse("reisen:reise", args=[self.r.pk]))
        self.assertContains(seite, "Kraftstoff: 88 L (Ø 11,1 L/100 km, 1 von 2 Etappen mit Verbrauchskurve TomTom)")
        ws = load_workbook(BytesIO(excel(self.reise(), stand_jetzt())))["Etappen"]
        self.assertEqual((ws["J2"].value, ws["K2"].value), ("TomTom", 65.4))
        self.assertEqual(ws["K3"].value, '=IF(OR(I3="",J3=""),0,I3*J3/100)')
